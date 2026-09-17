# -*- coding: utf-8 -*-
#!/usr/bin/env python
"""
hive-comment-update-from-mysql.py

从 MySQL 导出的 mysql_comment.csv（mysql_comment 页）读取字段注释，
与「大数据三期原始层及交换共享迁移进度-注释迁移.xlsx」终版页映射后，
将注释更新到 Hive ODS 表对应字段。

文件列结构（与截图一致）：
  CSV：用户名 | 表英文名 | 表中文名(不参与匹配) | 字段名 | 字段注释
  Excel 终版：二期-中心库账号(B) | 二期-源表(N) | 三期-hive-ods库 | 三期-ods表英文名(T) | 字段名(可选)

匹配规则（Excel 驱动，提高效率）：
  1. CSV 仅构建索引 {(用户名, 表英文名, 字段名): 注释}
  2. 仅遍历 Excel 终版页「迁移标识」=「是」的行，按二期账号+源表去 CSV 索引查找
  3. 命中后按 Hive 库+表（三期-hive-ods库、三期-ods表英文名）汇总分组，同表只处理一次
  4. 执行阶段按 Excel「三期-hive-ods库」分组：同库表连续处理，共用一条 PyHive 连接
  每组：一次元数据查询（SHOW TBLPROPERTIES + DESCRIBE）-> 分批 DDL（表注释 + 字段注释）
  若 Hive 目标表/字段已有注释：以 mysql_comment.csv 为准，合并该源表在 CSV 中的全部字段后更新
  若 Hive 尚无注释：仅当与 CSV 不一致时更新（减少无效 DDL）
  CSV 字段注释为空时：将 Hive 对应字段注释更新为空（COMMENT ''）
  单条 DDL 失败不中断任务，结束时汇总输出失败的表/字段注释明细
  Hive 目标表不存在：跳过该表（不计入失败），结束时汇总输出跳过明细

Hive 连接仅 PyHive（内联加载 lgbs-data-insert-kingsoft-prod-all.py 建连与 Kerberos）；按库复用会话。

用法：
  python hive-comment-update-from-mysql.py [--dry-run]
  python hive-comment-update-from-mysql.py <mysql_comment.csv> <mapping.xlsx> [--dry-run]

默认数据目录 /tmp/lgbs（与 user.txt 同目录）：
  /tmp/lgbs/mysql_comment.csv  或  /tmp/lgbs/mysql_comment.xlsx（mysql_comment 页）
  /tmp/lgbs/大数据三期原始层及交换共享迁移进度-注释迁移.xlsx

环境变量（可选）：
  HIVE_USER_CONFIG_FILE   Hive 连接配置（默认 /tmp/lgbs/user.txt）
  MYSQL_COMMENT_CSV       覆盖 MySQL 注释文件路径（csv/xlsx 均可）
  MYSQL_COMMENT_SHEET     xlsx 时的工作表名（默认：mysql_comment）
  MYSQL_COMMENT_ENCODING  强制文本编码（如 gb18030，MySQL 导出中文 CSV 常用）
  MIGRATION_XLSX          覆盖默认 Excel 路径
  XLSX_SHEET_NAME         Excel 工作表名（默认：终版）
  HIVE_ODS_DATABASE       Excel 无「三期-hive-ods库」列时的默认 Hive 库名
  HIVE_DDL_BATCH_SIZE     每张表 DDL 分批条数（默认 20）
  HIVE_KRB_REALM          Kerberos realm（/etc/krb5.conf 为 EXAMPLE.COM 时需设，如 LGXC_HADOOP.COM）
  HIVE_PYHIVE_TCP_SASL_SPLIT  IPv4 建连 + FQDN SASL（见 lgbs 脚本，no serverFQDN 时自动尝试）
  LGBS_PYHIVE_SCRIPT      lgbs-data-insert-kingsoft-prod-all.py 绝对路径（调度仅拷贝本脚本时必填）
  kinit_keytab / keytab   user.txt 中配置，供 PyHive 进程内 kinit（与 lgbs 一致）
  HIVE_COMMENT_CSV_WHEN_HIVE_HAS  Hive 已有注释时以 CSV 强制覆盖（默认 1）
  HIVE_COMMENT_FORCE_REAPPLY      与上同义保留；设为 0 可恢复「相同则跳过」
  HIVE_COMMENT_STRICT_EXIT        设为 0 时部分失败也返回 0（调度默认 1 有失败则 exit 1）
  DRY_RUN=1               仅生成 SQL、不执行
"""

from __future__ import print_function

import sys
import os
import re
import csv
import time
import codecs
import importlib.util
import zipfile
import xml.etree.ElementTree as ET

try:
    text_type = unicode  # noqa: F821 (py2)
except Exception:
    text_type = str


def _raise_csv_field_size_limit():
    """
    MySQL 注释 CSV 单字段可能很长；默认 128KB 限制会导致整文件读取失败。
    """
    try:
        limit = sys.maxsize
    except Exception:
        limit = 2147483647
    while True:
        try:
            csv.field_size_limit(int(limit))
            return
        except OverflowError:
            limit = int(limit / 10)
        except Exception:
            try:
                csv.field_size_limit(10 * 1024 * 1024)
            except Exception:
                pass
            return


_raise_csv_field_size_limit()


def to_text(val):
    if val is None:
        return text_type("")
    try:
        if isinstance(val, bytes):
            for enc in ("utf-8", "gb18030", "gbk", "latin-1"):
                try:
                    return val.decode(enc, "replace")
                except Exception:
                    continue
            return val.decode("utf-8", "replace")
    except Exception:
        pass
    try:
        return text_type(val)
    except Exception:
        try:
            b = str(val)
            if isinstance(b, bytes):
                return text_type(b.decode("utf-8", "replace"))
            return text_type(b)
        except Exception:
            return text_type("")


def _print_u(msg):
    try:
        if isinstance(msg, text_type):
            sys.stdout.write(msg.encode("utf-8") + "\n")
        else:
            sys.stdout.write(to_text(msg).encode("utf-8") + "\n")
    except Exception:
        try:
            print(msg)
        except Exception:
            pass


def _now():
    return time.strftime("%Y-%m-%d %H:%M:%S")


# ==================== PyHive 建连（内联，避免调度环境缺少 hive_pyhive_conn 模块） ====================
_HPC_LGBS_MOD = None


def _hpc_lgbs_script_candidates():
    """查找 lgbs-data-insert-kingsoft-prod-all.py（调度常只拷贝本脚本到 /tmp）。"""
    name = "lgbs-data-insert-kingsoft-prod-all.py"
    out = []
    seen = set()

    def _add(p):
        try:
            p = os.path.abspath(to_text(p))
        except Exception:
            return
        if p in seen or not os.path.isfile(p):
            return
        seen.add(p)
        out.append(p)

    for env_k in ("LGBS_PYHIVE_SCRIPT", "HIVE_PYHIVE_LGBS_SCRIPT", "LGBS_SCRIPT_PATH"):
        ep = (os.getenv(env_k, "") or "").strip()
        if ep:
            _add(ep)

    try:
        here = os.path.dirname(os.path.abspath(__file__))
        _add(os.path.join(here, name))
    except Exception:
        pass

    try:
        _add(os.path.join(os.getcwd(), name))
    except Exception:
        pass

    data_dir = (os.getenv("LGBS_DATA_DIR", "") or "/tmp/lgbs").strip() or "/tmp/lgbs"
    _add(os.path.join(data_dir, name))

    for base in (
        os.getenv("KINGSOFT_API_HOME", ""),
        os.getenv("LGBS_SCRIPT_DIR", ""),
        "/opt/kingsoft-api",
        "/home/ods/kingsoft-api",
    ):
        b = to_text(base).strip()
        if b:
            _add(os.path.join(b, name))
    return out


def _hpc_get_lgbs():
    global _HPC_LGBS_MOD
    if _HPC_LGBS_MOD is not None:
        return _HPC_LGBS_MOD
    candidates = _hpc_lgbs_script_candidates()
    if not candidates:
        raise ImportError(
            u"未找到 lgbs-data-insert-kingsoft-prod-all.py。"
            u"请设置 LGBS_PYHIVE_SCRIPT=绝对路径，或将 lgbs 脚本放到 /tmp/lgbs/ 或与本文同目录。"
        )
    path = candidates[0]
    try:
        spec = importlib.util.spec_from_file_location("lgbs_pyhive_shim", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
    except Exception:
        try:
            import imp  # py2

            mod = imp.load_source("lgbs_pyhive_shim", path)
        except Exception as ex:
            raise ImportError(
                u"加载 {0} 失败：{1}".format(path, to_text(ex))
            )
    mod._print_u = _print_u
    try:
        mod._now = _now
    except Exception:
        pass
    _HPC_LGBS_MOD = mod
    return _HPC_LGBS_MOD


def _hpc_pyhive_import_available():
    try:
        return bool(_hpc_get_lgbs()._hive_pyhive_import_available())
    except Exception:
        return False


def _hpc_last_pyhive_import_error():
    try:
        return _hpc_get_lgbs()._LAST_PYHIVE_IMPORT_ERROR or ""
    except Exception:
        return ""


def _hpc_open_pyhive_connection(hive_cfg, database=None):
    lg = _hpc_get_lgbs()
    if database:
        hc = dict(hive_cfg or {})
        hc["database"] = database
        return lg._open_pyhive_connection(hc)
    return lg._open_pyhive_connection(hive_cfg)


def _trim_token(val):
    if val is None:
        return ""
    try:
        t = to_text(val)
        return t.strip().strip(u"\u3000")
    except Exception:
        return str(val).strip()


def _normalize_identifier_token(val):
    t = _trim_token(val)
    if not t:
        return ""
    t = t.replace(u"\ufeff", u"").replace(u"\u200b", u"").replace(u"\u200c", u"").replace(u"\u200d", u"")
    try:
        t = re.sub(u"[\x00-\x1f\x7f]", u"", t)
    except Exception:
        pass
    return _trim_token(t)


def _normalize_match_key(val):
    """表/用户匹配键：去空白、统一小写。"""
    t = _normalize_identifier_token(val)
    return t.lower() if t else ""


def _flag_is_yes(val):
    """迁移标识栏位是否为「是」，兼容常见写法。"""
    if val is None:
        return False
    t = _trim_token(val)
    if not t:
        return False
    u = t.lower()
    return t == u"是" or u in ("y", "yes", "true", "1")


def _escape_hive_comment(s):
    if s is None:
        return ""
    return to_text(s).replace("'", "''")


def _safe_hive_ident(name):
    """库名等规范化（小写）。"""
    t = to_text(name).strip()
    if not t:
        return ""
    t2 = re.sub(r"[^0-9a-zA-Z_]", "_", t)
    if re.match(r"^\d", t2):
        t2 = "_" + t2
    return t2.lower()


def _quote_hive_ident(name):
    """DDL/DESCRIBE 用：保留 Hive 原始列名/表名，仅转义反引号。"""
    t = to_text(name).strip()
    if not t:
        return ""
    return t.replace("`", "``")


_HIVE_KRB_PLACEHOLDER_REALMS = frozenset(
    ("EXAMPLE.COM", "EXAMPLE.ORG", "INVALID.REALM", "TEST.REALM", "MYREALM.ORG")
)
_HIVE_KRB_REALM_FALLBACK_LOGGED = False


def _hive_kerberos_realm_is_placeholder(realm):
    t = to_text(realm or u"").strip().upper()
    return bool(t and t in _HIVE_KRB_PLACEHOLDER_REALMS)


def _hive_guess_kerberos_realm_from_krbhost(krbhost):
    h = to_text(krbhost or u"").strip()
    if not h or re.match(r"^(?:\d{1,3}\.){3}\d{1,3}$", h) or "." not in h:
        return ""
    parts = h.lower().split(".")
    if len(parts) >= 3 and parts[0] in ("hadoop", "hive", "hs2", "nn", "rm", "nm"):
        dom = ".".join(parts[1:])
    else:
        dom = ".".join(parts)
    return dom.upper()


def _hive_resolve_kerberos_realm(hive_cfg):
    global _HIVE_KRB_REALM_FALLBACK_LOGGED
    allow_ph = _env_truthy("HIVE_KRB_ALLOW_PLACEHOLDER_REALM", False)

    def _usable(r):
        t = to_text(r or u"").strip()
        if not t:
            return ""
        if _hive_kerberos_realm_is_placeholder(t) and not allow_ph:
            return ""
        return t

    env_r = _usable((os.getenv("HIVE_KRB_REALM", "") or "").strip())
    if env_r:
        return env_r
    krb5_r = _usable(_hive_read_default_realm_from_krb5_conf())
    if krb5_r:
        return krb5_r
    guessed = _usable(_hive_guess_kerberos_realm_from_krbhost((hive_cfg or {}).get("krbhost")))
    if guessed and not _HIVE_KRB_REALM_FALLBACK_LOGGED:
        _HIVE_KRB_REALM_FALLBACK_LOGGED = True
        _print_u(
            u"【{0}】Kerberos realm 从 krbhost 推断为 {1}（可设 HIVE_KRB_REALM 覆盖）".format(
                _now(), guessed
            )
        )
    return guessed


def _hive_read_default_realm_from_krb5_conf():
    for fp in (
        (os.getenv("KRB5_CONFIG", "") or "").strip(),
        "/etc/krb5.conf",
    ):
        if not fp or not os.path.isfile(fp):
            continue
        try:
            with open(fp, "r") as _f:
                for _ln in _f:
                    _ln2 = _ln.strip()
                    if not _ln2 or _ln2.startswith("#") or _ln2.startswith(";"):
                        continue
                    if "default_realm" in _ln2 and "=" in _ln2:
                        _k, _v = _ln2.split("=", 1)
                        if _k.strip().lower() == "default_realm":
                            return _v.strip()
        except Exception:
            continue
    return ""


# ==================== Excel 列名（终版页） ====================
F_CENTER = u"二期-中心库账号"
F_SOURCE = u"二期-源表"
F_HIVE_DB = u"三期-hive-ods库"
F_HIVE_TABLE = u"三期-ods表英文名"
F_FIELD = u"字段名"
F_MIGRATE_FLAG = u"迁移标识"

# ==================== CSV 列名（mysql_comment 页导出） ====================
C_USER = u"用户名"
C_TABLE = u"表英文名"
C_TABLE_CN = u"表中文名"
C_FIELD = u"字段名"
C_COMMENT = u"字段注释"

_CELL_REF_RE = re.compile(r"^([A-Z]+)(\d+)$", re.I)


# ==================== 默认路径（生产环境 /tmp/lgbs） ====================
LGBS_DATA_DIR = "/tmp/lgbs"
DEFAULT_MYSQL_COMMENT_CSV = os.path.join(LGBS_DATA_DIR, "mysql_comment.csv")
DEFAULT_MYSQL_COMMENT_XLSX = os.path.join(LGBS_DATA_DIR, "mysql_comment.xlsx")
DEFAULT_MIGRATION_XLSX = os.path.join(
    LGBS_DATA_DIR, u"大数据三期原始层及交换共享迁移进度-注释迁移.xlsx"
)

# ==================== Hive 配置 / 执行（对齐 oracle-table-structure-move-hive-prod.py） ====================
HIVE_USER_CONFIG_FILE = os.path.join(LGBS_DATA_DIR, "user.txt")
_LAST_HIVE_CFG_FILE = None
_HIVE_CFG_LOGGED = False


def _hive_user_txt_line_is_ini_section_header(ln):
    t = to_text(ln).strip()
    return len(t) >= 3 and t.startswith("[") and t.endswith("]")


def read_hive_user_config(file_path=HIVE_USER_CONFIG_FILE):
    cfg = {}
    fp = file_path
    fp = os.getenv("HIVE_USER_CONFIG_FILE", fp) or fp
    if not os.path.isabs(fp):
        local_fp = os.path.join(os.path.dirname(os.path.abspath(__file__)), fp)
        if os.path.isfile(local_fp):
            fp = local_fp
    global _LAST_HIVE_CFG_FILE
    _LAST_HIVE_CFG_FILE = fp
    f = None
    try:
        if sys.version_info[0] < 3:
            f = codecs.open(fp, "r", encoding="utf-8")
        else:
            f = open(fp, "r", encoding="utf-8")
        for line in f:
            ln = (line or "").strip()
            if not ln or ln.startswith("#"):
                continue
            if _hive_user_txt_line_is_ini_section_header(ln):
                break
            if "=" in ln:
                k, v = ln.split("=", 1)
                kk = k.strip()
                try:
                    if kk and isinstance(kk, text_type) and kk[0] == u"\ufeff":
                        kk = kk.lstrip(u"\ufeff")
                except Exception:
                    pass
                try:
                    kk2 = to_text(kk).strip().lower()
                except Exception:
                    kk2 = str(kk).strip().lower()
                cfg[kk2] = (v.strip() if hasattr(v, "strip") else to_text(v).strip())
    except IOError as e:
        raise RuntimeError("Hive 配置文件不存在或不可读：{0}，err={1}".format(fp, e))
    except Exception as e:
        raise RuntimeError("读取 Hive 配置文件失败：{0} err={1}".format(fp, e))
    finally:
        try:
            if f:
                f.close()
        except Exception:
            pass
    return cfg


def get_hive_config(database):
    """user.txt 的 st= 映射为 host；kinit_keytab 等与 lgbs 脚本一致。"""
    cfg = read_hive_user_config()
    host = (cfg.get("st") or cfg.get("host") or os.getenv("HIVE_HOST", "") or "").strip()
    port_raw = cfg.get("port", "")
    try:
        port = int(port_raw) if str(port_raw).isdigit() else 21066
    except Exception:
        port = 21066
    _ar = cfg.get("auth")
    if _ar is None or not to_text(_ar).strip():
        _auth_m = u"KERBEROS"
    else:
        _auth_m = to_text(_ar).strip()
    global _HIVE_CFG_LOGGED
    if not _HIVE_CFG_LOGGED:
        _HIVE_CFG_LOGGED = True
        _print_u(
            u"【{0}】Hive 配置：file={1}，host={2}，port={3}，jdbc_url_set={4}".format(
                to_text(_now()),
                to_text(_LAST_HIVE_CFG_FILE),
                host or u"(空)",
                port,
                bool((cfg or {}).get("jdbc_url")),
            )
        )
    if not host and not (cfg.get("jdbc_url") or "").strip():
        raise RuntimeError(
            u"Hive host 为空（failed to resolve sockaddr for :{0}），"
            u"请在 {1} 配置 st=IP 或 jdbc_url=（须在 [libdefaults] 等 krb5 段之前）".format(
                port, to_text(_LAST_HIVE_CFG_FILE)
            )
        )
    return {
        "host": host,
        "port": port,
        "username": cfg.get("username", ""),
        "database": database or cfg.get("database", "") or "default",
        "auth": _auth_m,
        "kerberos_service_name": cfg.get("kerberos_service_name", "hive"),
        "krbhost": cfg.get("krbhost", ""),
        "jdbc_url": cfg.get("jdbc_url", "") or "",
        "kinit_keytab": to_text(
            cfg.get("kinit_keytab", "")
            or cfg.get("keytab", "")
            or cfg.get("kerberos_keytab", "")
            or ""
        ).strip(),
        "kinit_user": to_text(cfg.get("kinit_user", "") or "").strip(),
    }


def _hive_describe_pyhive(cur, hive_db, hive_table, skip_use=False):
    db_ident = _safe_hive_ident(hive_db)
    tb_ident = _quote_hive_ident(hive_table)
    if not skip_use:
        cur.execute(u"USE `{0}`".format(db_ident))
    cur.execute(u"DESCRIBE `{0}`".format(tb_ident))
    rows = cur.fetchall() or []
    cols = {}
    for r in rows:
        if not r:
            continue
        cn = to_text(r[0]).strip() if r[0] is not None else ""
        dt = to_text(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
        cm = _trim_token(r[2]) if len(r) > 2 else ""
        if not cn:
            continue
        if cn.lower() == "col_name" and dt.lower() == "data_type":
            continue
        if cn.startswith("#"):
            break
        cols[_normalize_match_key(cn)] = {"name": cn, "type": dt, "comment": cm}
    return cols


def _hive_sql_strip_trailing_semicolons(sql_text):
    """PyHive/HiveServer2 不接受语句末尾分号（与 lgbs 脚本一致）。"""
    s = to_text(sql_text or u"").strip()
    while s.endswith(u";"):
        s = s[:-1].strip()
    return s


def build_use_db_sql(hive_db):
    db_ident = _safe_hive_ident(hive_db)
    return u"USE `{0}`".format(db_ident)


def build_alter_table_comment_stmt(hive_table, new_comment):
    tb_ident = _quote_hive_ident(hive_table)
    cmt = _escape_hive_comment(new_comment)
    return u"ALTER TABLE `{0}` SET TBLPROPERTIES ('comment' = '{1}')".format(tb_ident, cmt)


def build_alter_column_comment_stmt(hive_table, col_name, col_type, new_comment):
    tb_ident = _quote_hive_ident(hive_table)
    col_ident = _quote_hive_ident(col_name)
    typ = to_text(col_type).strip() or "string"
    cmt = _escape_hive_comment(new_comment)
    return (
        u"ALTER TABLE `{0}` CHANGE COLUMN `{1}` `{1}` {2} COMMENT '{3}'"
    ).format(tb_ident, col_ident, typ, cmt)


def build_table_write_batch_sql(hive_db, alter_stmts):
    """同一张表的多条 ALTER；返回语句列表（USE 仅一次在首位，逐条 cur.execute）。"""
    stmts = [to_text(s).strip() for s in (alter_stmts or []) if to_text(s).strip()]
    if not stmts:
        return []
    return [build_use_db_sql(hive_db)] + stmts


def _format_sql_batch_for_log(sql_batch):
    if isinstance(sql_batch, (list, tuple)):
        return u"\n".join(
            to_text(s).strip() for s in sql_batch if to_text(s).strip()
        )
    return to_text(sql_batch or u"")


def _sql_batch_is_empty(sql_batch):
    if sql_batch is None:
        return True
    if isinstance(sql_batch, (list, tuple)):
        return not any(to_text(s).strip() for s in sql_batch)
    return not to_text(sql_batch).strip()


def _hive_table_type_from_props(props):
    if not props:
        return ""
    for k in ("table type", "table_type"):
        if k in props:
            return to_text(props[k]).strip().upper()
    return ""


def _hive_is_virtual_view(table_type, hive_table):
    """仅当元数据 Table Type=VIRTUAL_VIEW 时视为视图（不按表名 _view 猜测）。"""
    tt = to_text(table_type or "").strip().upper()
    return tt in ("VIRTUAL_VIEW", "MATERIALIZED_VIEW")


def _hive_fetch_table_metadata_pyhive(cur, hive_db, hive_table):
    """PyHive 同一会话内读取表元数据（USE 一次）。"""
    db_ident = _safe_hive_ident(hive_db)
    tb_ident = _quote_hive_ident(hive_table)
    cur.execute(u"USE `{0}`".format(db_ident))
    table_cmt = u""
    table_type = u""
    try:
        cur.execute(u"SHOW TBLPROPERTIES `{0}`".format(tb_ident))
        props = {}
        for row in cur.fetchall() or []:
            if not row or len(row) < 2:
                continue
            k = to_text(row[0]).strip().lower()
            v = _trim_token(row[1])
            props[k] = v
            if k == "comment":
                table_cmt = v
        table_type = _hive_table_type_from_props(props)
    except Exception:
        pass
    cols = _hive_describe_pyhive(cur, hive_db, hive_table, skip_use=True)
    return cols, table_cmt, table_type


def _is_hive_session_down_error(err):
    t = to_text(err).lower()
    return (
        "session is down" in t
        or "session closed" in t
        or "invalid session" in t
        or "session expired" in t
    )


def _hive_ddl_batch_size():
    raw = (os.getenv("HIVE_DDL_BATCH_SIZE", "") or "20").strip()
    try:
        return max(1, int(raw))
    except Exception:
        return 20


def _env_truthy(name, default=False):
    v = (os.getenv(name, "") or "").strip().lower()
    if not v:
        return default
    return v in ("1", "true", "yes", "on", "y")


def _close_pyhive(conn, cur):
    for obj in (cur, conn):
        if obj is None:
            continue
        try:
            obj.close()
        except Exception:
            pass


def _split_hive_sql_statements(sql_text):
    """
    按分号拆分 SQL，忽略单引号字符串内的分号（COMMENT '...;...' 不能 naive split）。
    支持 Hive 转义 ''。
    """
    s = to_text(sql_text or u"")
    out = []
    buf = []
    in_quote = False
    i = 0
    n = len(s)
    while i < n:
        ch = s[i]
        if ch == u"'":
            if in_quote and i + 1 < n and s[i + 1] == u"'":
                buf.append(u"''")
                i += 2
                continue
            in_quote = not in_quote
            buf.append(ch)
        elif ch == u";" and not in_quote:
            st = u"".join(buf).strip()
            if st:
                out.append(st)
            buf = []
        else:
            buf.append(ch)
        i += 1
    st = u"".join(buf).strip()
    if st:
        out.append(st)
    return out


def _iter_hive_sql_statements(sql_batch):
    """展开为单条语句；支持 list 或旧版换行拼接字符串。"""
    if isinstance(sql_batch, (list, tuple)):
        for item in sql_batch:
            st = _hive_sql_strip_trailing_semicolons(item)
            if st:
                yield st
        return
    s = to_text(sql_batch or u"").strip()
    if not s:
        return
    parts = _split_hive_sql_statements(s)
    if len(parts) == 1 and u"\n" in parts[0]:
        for line in parts[0].splitlines():
            st = _hive_sql_strip_trailing_semicolons(line)
            if st:
                yield st
        return
    for stmt in parts:
        st = _hive_sql_strip_trailing_semicolons(stmt)
        if st:
            yield st


def _pyhive_run_sql_on_cursor(cur, sql_batch):
    for stmt in _iter_hive_sql_statements(sql_batch):
        cur.execute(stmt)


def _is_hive_table_not_found_error(err):
    t = to_text(err).lower()
    return (
        "table not found" in t
        or "error 10001" in t
        or "42s02" in t
    )


def _hive_table_name_candidates(hive_db, hive_table):
    """Excel 表名与 Hive 实际表名不一致时尝试多种写法。"""
    t = to_text(hive_table).strip()
    db = (_safe_hive_ident(hive_db) or "").strip()
    seen = set()
    out = []

    def _add(name):
        n = to_text(name).strip()
        if not n or n in seen:
            return
        seen.add(n)
        out.append(n)

    if "." in t:
        parts = [p for p in t.split(".") if p]
        if len(parts) >= 2:
            _add(parts[-1])
            _add("_".join(parts))
    _add(t)
    if db:
        pref = db + "_"
        tl = t.lower()
        pl = pref.lower()
        if tl.startswith(pl) and len(t) > len(pref):
            _add(t[len(pref) :])
        elif not tl.startswith(pl):
            _add(pref + t)
    return out


def _hive_fetch_table_metadata_pyhive_resolved(cur, hive_db, hive_table):
    """
    读取表元数据；Table not found 时按候选表名重试。
    返回 (cols, table_cmt, table_type, resolved_table_name)。
    """
    last_err = None
    cands = _hive_table_name_candidates(hive_db, hive_table)
    for cand in cands:
        try:
            cols, table_cmt, table_type = _hive_fetch_table_metadata_pyhive(
                cur, hive_db, cand
            )
            if cand != hive_table:
                _print_u(
                    u"【{0}】Hive 表名解析：{1}.{2} -> {1}.{3}".format(
                        _now(), hive_db, hive_table, cand
                    )
                )
            return cols, table_cmt, table_type, cand
        except Exception as ex:
            last_err = ex
            if _is_hive_table_not_found_error(ex):
                continue
            raise
    if last_err:
        raise last_err
    raise RuntimeError(
        u"Hive 表不存在：{0}.{1}（已尝试 {2}）".format(
            hive_db, hive_table, u", ".join(cands[:5])
        )
    )


class _PyHiveSessionPool(object):
    """按 Hive 库复用 PyHive 连接（lgbs Kerberos 逻辑）。"""

    def __init__(self):
        self.enabled = False
        self._sessions = {}

    def acquire(self, hive_db):
        db = _safe_hive_ident(hive_db) or "default"
        if db not in self._sessions:
            hive_cfg = get_hive_config(db)
            conn = _hpc_open_pyhive_connection(hive_cfg, database=db)
            cur = conn.cursor()
            self._sessions[db] = (conn, cur)
        return self._sessions[db]

    def reconnect(self, hive_db):
        self.release(hive_db)
        return self.acquire(hive_db)

    def release(self, hive_db):
        db = _safe_hive_ident(hive_db) or "default"
        pair = self._sessions.pop(db, None)
        if pair:
            _close_pyhive(pair[0], pair[1])

    def close_all(self):
        for db in list(self._sessions.keys()):
            self.release(db)


def _pyhive_execute_sql(hive_cfg, sql_batch, max_retry=2, cur=None, hive_db=None):
    """
    执行 SQL；若传入 cur 则复用会话，否则新建短连接。
    session is down 时自动重连重试（仅短连接模式）。
    """
    if cur is not None:
        _pyhive_run_sql_on_cursor(cur, sql_batch)
        return True

    last_err = None
    db = hive_db or (hive_cfg or {}).get("database") or "default"
    for attempt in range(max_retry):
        conn = None
        cursor = None
        try:
            hive_cfg = get_hive_config(db)
            conn = _hpc_open_pyhive_connection(hive_cfg, database=db)
            cursor = conn.cursor()
            _pyhive_run_sql_on_cursor(cursor, sql_batch)
            return True
        except Exception as ex:
            last_err = ex
            if _is_hive_session_down_error(ex) and attempt < max_retry - 1:
                _print_u(
                    u"【{0}】WARNING: Hive session 断开，重连重试 ({1}/{2})".format(
                        _now(), attempt + 2, max_retry
                    )
                )
                continue
            raise
        finally:
            _close_pyhive(conn, cursor)
    if last_err:
        raise last_err
    return True


def _execute_hive_sql_batch(
    sql_batch,
    hive_cfg,
    dry_run,
    label,
    pyhive_cur=None,
    pyhive_pool=None,
    hive_db=None,
):
    """执行一批 SQL（PyHive），返回 (是否成功, 错误信息)。"""
    if _sql_batch_is_empty(sql_batch):
        return True, u""
    if dry_run:
        _print_u(
            u"【DRY-RUN】{0}\n{1}".format(label, _format_sql_batch_for_log(sql_batch))
        )
        return True, u""
    try:
        db = hive_db or (hive_cfg or {}).get("database") or "default"
        try:
            if pyhive_cur is not None:
                _pyhive_run_sql_on_cursor(pyhive_cur, sql_batch)
            else:
                _pyhive_execute_sql(hive_cfg, sql_batch, hive_db=db)
        except Exception as ex:
            if pyhive_pool and pyhive_pool.enabled and _is_hive_session_down_error(ex):
                _print_u(
                    u"【{0}】WARNING: PyHive session 断开，按库重连后重试：{1}".format(
                        _now(), label
                    )
                )
                _, pyhive_cur = pyhive_pool.reconnect(db)
                _pyhive_run_sql_on_cursor(pyhive_cur, sql_batch)
            else:
                raise
        _print_u(u"【{0}】OK {1}".format(_now(), label))
        return True, u""
    except Exception as ex:
        err_msg = to_text(ex)[:2000]
        _print_u(u"【{0}】FAIL {1} err={2}".format(_now(), label, err_msg))
        return False, err_msg


def _parse_alter_comment_stmt(stmt):
    """从 ALTER 语句解析表/列注释信息，供失败汇总使用。"""
    s = to_text(stmt or u"").strip()
    info = {"kind": "column", "column": u"", "comment": u""}
    if not s:
        return info
    up = s.upper()
    if "SET TBLPROPERTIES" in up:
        info["kind"] = "table"
        m = re.search(r"comment'\s*=\s*'((?:[^']|'')*)'", s, re.I)
        if m:
            info["comment"] = m.group(1).replace("''", "'")
        return info
    m_col = re.search(r"CHANGE\s+COLUMN\s+`([^`]+)`", s, re.I)
    if m_col:
        info["column"] = m_col.group(1)
    m_cmt = re.search(r"COMMENT\s+'((?:[^']|'')*)'", s, re.I)
    if m_cmt:
        info["comment"] = m_cmt.group(1).replace("''", "'")
    return info


def _make_fail_record(full_table, kind, column, comment, stmt, error):
    return {
        "full_table": full_table,
        "kind": kind,
        "column": to_text(column or u""),
        "comment": _trim_token(comment),
        "stmt": to_text(stmt or u"")[:500],
        "error": to_text(error or u"")[:2000],
    }


def _make_missing_table_skip_record(full_table, table_comment, field_count, error):
    return {
        "full_table": full_table,
        "table_comment": _trim_token(table_comment),
        "field_count": max(0, int(field_count or 0)),
        "error": to_text(error or u"")[:500],
    }


def _print_missing_tables_summary(skip_records):
    """汇总输出 Hive 中不存在的映射表（已跳过，不计入失败）。"""
    if not skip_records:
        return
    field_total = sum(int(r.get("field_count") or 0) for r in skip_records)
    print("=" * 80)
    _print_u(
        u"【{0}】Hive 表不存在跳过汇总（共 {1} 张，涉及字段计划 {2} 条）".format(
            _now(), len(skip_records), field_total
        )
    )
    for i, r in enumerate(skip_records, 1):
        ft = r.get("full_table") or u""
        cmt = r.get("table_comment") or u""
        if len(cmt) > 80:
            cmt = cmt[:80] + u"..."
        _print_u(
            u"  [{0}] {1} | 表注释={2} | 字段计划={3}".format(
                i, ft, cmt or u"(无)", r.get("field_count", 0)
            )
        )
    print("=" * 80)


def _print_failures_summary(fail_records):
    """汇总输出执行失败的表注释、字段注释。"""
    if not fail_records:
        return
    table_fails = [r for r in fail_records if r.get("kind") in ("table", "metadata")]
    col_fails = [r for r in fail_records if r.get("kind") == "column"]

    print("=" * 80)
    _print_u(
        u"【{0}】失败明细汇总（共 {1} 条：表注释/元数据 {2}，字段注释 {3}）".format(
            _now(), len(fail_records), len(table_fails), len(col_fails)
        )
    )
    if table_fails:
        _print_u(u"  ---------- 表注释失败 ----------")
        for i, r in enumerate(table_fails, 1):
            ft = r.get("full_table") or u""
            cmt = r.get("comment") or u""
            if len(cmt) > 80:
                cmt = cmt[:80] + u"..."
            _print_u(
                u"  [{0}] {1} | 注释={2}".format(i, ft, cmt or u"(无)")
            )
            if r.get("kind") == "metadata":
                _print_u(u"       阶段=读取元数据")
            if r.get("error"):
                _print_u(u"       错误={0}".format(r["error"][:500]))
            if r.get("stmt"):
                _print_u(u"       DDL={0}".format(r["stmt"][:300]))
    if col_fails:
        _print_u(u"  ---------- 字段注释失败 ----------")
        for i, r in enumerate(col_fails, 1):
            ft = r.get("full_table") or u""
            col = r.get("column") or u"?"
            cmt = r.get("comment") or u""
            if len(cmt) > 60:
                cmt = cmt[:60] + u"..."
            _print_u(
                u"  [{0}] {1}.{2} | 注释={3}".format(i, ft, col, cmt or u"(无)")
            )
            if r.get("error"):
                _print_u(u"       错误={0}".format(r["error"][:500]))
    print("=" * 80)


def _execute_alter_stmts_chunked(
    hive_db,
    alter_stmts,
    hive_cfg,
    dry_run,
    full_table,
    n_table_stmts=0,
    pyhive_cur=None,
    pyhive_pool=None,
):
    """
    将 ALTER 分批执行；某批失败时对该批内语句逐条重试，单条失败不中断同表其余语句。
    返回 (ok_table, fail_table, ok_col, fail_col, failures)。
    """
    failures = []
    if not alter_stmts:
        return 0, 0, 0, 0, failures
    batch_size = _hive_ddl_batch_size()
    n_table = max(0, min(int(n_table_stmts or 0), len(alter_stmts)))
    ok_table = fail_table = ok_col = fail_col = 0

    def _bump_ok(stmt_idx):
        nonlocal ok_table, ok_col
        if stmt_idx < n_table:
            ok_table += 1
        else:
            ok_col += 1

    def _bump_fail(stmt_idx, stmt, err_msg):
        nonlocal fail_table, fail_col
        parsed = _parse_alter_comment_stmt(stmt)
        kind = "table" if stmt_idx < n_table else parsed.get("kind", "column")
        failures.append(
            _make_fail_record(
                full_table,
                kind,
                parsed.get("column", u""),
                parsed.get("comment", u""),
                stmt,
                err_msg,
            )
        )
        if stmt_idx < n_table:
            fail_table += 1
        else:
            fail_col += 1

    chunks = [
        (i, alter_stmts[i : i + batch_size])
        for i in range(0, len(alter_stmts), batch_size)
    ]
    for chunk_start, chunk in chunks:
        sql_batch = build_table_write_batch_sql(hive_db, chunk)
        label = u"{0} [ddl batch @{1}, stmts={2}]".format(
            full_table, chunk_start + 1, len(chunk)
        )
        ok, err_msg = _execute_hive_sql_batch(
            sql_batch,
            hive_cfg,
            dry_run,
            label,
            pyhive_cur=pyhive_cur,
            pyhive_pool=pyhive_pool,
            hive_db=hive_db,
        )
        if ok:
            for j in range(len(chunk)):
                _bump_ok(chunk_start + j)
            continue
        if len(chunk) == 1:
            _bump_fail(chunk_start, chunk[0], err_msg)
            continue
        _print_u(
            u"【{0}】批次失败，逐条重试：{1}（{2} 条）".format(
                _now(), full_table, len(chunk)
            )
        )
        for j, stmt in enumerate(chunk):
            one_sql = build_table_write_batch_sql(hive_db, [stmt])
            one_label = u"{0} [ddl {1}/{2}]".format(
                full_table, chunk_start + j + 1, len(alter_stmts)
            )
            ok_one, err_one = _execute_hive_sql_batch(
                one_sql,
                hive_cfg,
                dry_run,
                one_label,
                pyhive_cur=pyhive_cur,
                pyhive_pool=pyhive_pool,
                hive_db=hive_db,
            )
            if ok_one:
                _bump_ok(chunk_start + j)
            else:
                _bump_fail(chunk_start + j, stmt, err_one)
                _print_u(
                    u"【{0}】FAIL 单条 DDL：{1}".format(_now(), to_text(stmt)[:500])
                )
    return ok_table, fail_table, ok_col, fail_col, failures


# ==================== 数据加载 ====================
def _pick_header_index(headers, candidates):
    norm = {}
    for i, h in enumerate(headers):
        key = _normalize_match_key(h)
        if key:
            norm[key] = i
    for cand in candidates:
        ck = _normalize_match_key(cand)
        if ck in norm:
            return norm[ck]
    return None


def _row_get(row, idx):
    if idx is None or idx < 0 or idx >= len(row):
        return ""
    return row[idx]


def _file_is_xlsx(path):
    try:
        with open(path, "rb") as fb:
            return fb.read(2) == b"PK"
    except Exception:
        return False


def _comment_csv_encodings():
    forced = (os.getenv("MYSQL_COMMENT_ENCODING", "") or "").strip()
    if forced:
        return [forced]
    # MySQL/Windows 导出中文 CSV 多为 GBK 系；head 乱码而英文正常时优先 gb18030
    return ("gb18030", "gbk", "utf-8-sig", "utf-8", "cp936", "latin-1")


def _read_delimited_text_rows(file_path):
    """尝试多种编码与分隔符读取 csv/tsv（已放宽 csv.field_size_limit）。"""
    last_err = None
    for enc in _comment_csv_encodings():
        for delim in (",", "\t", ";"):
            rows = []
            f = None
            try:
                if sys.version_info[0] < 3:
                    f = codecs.open(file_path, "r", encoding=enc, errors="replace")
                    reader = csv.reader(f, delimiter=delim)
                else:
                    f = open(file_path, "r", encoding=enc, newline="", errors="replace")
                    reader = csv.reader(f, delimiter=delim)
                for row in reader:
                    rows.append([to_text(c) for c in row])
                if rows and any(any(_trim_token(c) for c in r) for r in rows):
                    return rows, enc, repr(delim)
            except Exception as e:
                last_err = e
                rows = []
            finally:
                try:
                    if f:
                        f.close()
                except Exception:
                    pass
    return [], last_err, ""


def _load_comment_source_rows(file_path):
    """
    加载 MySQL 注释源表：支持 .csv / .tsv / 实为 xlsx 的文件，以及 mysql_comment.xlsx。
    返回 (rows_raw, source_desc)
    """
    if not os.path.isfile(file_path):
        tried = [file_path]
        for alt in (
            DEFAULT_MYSQL_COMMENT_CSV,
            DEFAULT_MYSQL_COMMENT_XLSX,
            os.path.splitext(file_path)[0] + ".xlsx",
            os.path.splitext(file_path)[0] + ".csv",
        ):
            if alt and alt not in tried and os.path.isfile(alt):
                file_path = alt
                break
        else:
            raise RuntimeError(
                u"MySQL 注释文件不存在。已查找：{0}；请放置 {1} 或 {2}".format(
                    file_path, DEFAULT_MYSQL_COMMENT_CSV, DEFAULT_MYSQL_COMMENT_XLSX
                )
            )

    sheet_name = (os.getenv("MYSQL_COMMENT_SHEET", "") or "mysql_comment").strip() or "mysql_comment"

    if _file_is_xlsx(file_path) or to_text(file_path).lower().endswith(".xlsx"):
        try:
            matrix = _xlsx_read_sheet_rows_stdlib(file_path, sheet_name)
        except Exception as e:
            raise RuntimeError(
                u"读取 xlsx 失败：{0} sheet={1} err={2}".format(file_path, sheet_name, to_text(e))
            )
        if not matrix:
            raise RuntimeError(
                u"xlsx 工作表为空：{0} sheet={1}".format(file_path, sheet_name)
            )
        return matrix, u"xlsx:{0}".format(sheet_name)

    rows_raw, last_err, delim_info = _read_delimited_text_rows(file_path)
    if rows_raw:
        return rows_raw, u"text enc/delim={0}".format(delim_info)

    # csv 路径读不出时，尝试同目录 mysql_comment.xlsx
    xlsx_alt = os.path.join(os.path.dirname(file_path), "mysql_comment.xlsx")
    if not os.path.isfile(xlsx_alt):
        xlsx_alt = DEFAULT_MYSQL_COMMENT_XLSX
    if os.path.isfile(xlsx_alt):
        _print_u(
            u"【{0}】WARNING: {1} 无法按文本解析，改用 {2}".format(_now(), file_path, xlsx_alt)
        )
        matrix = _xlsx_read_sheet_rows_stdlib(xlsx_alt, sheet_name)
        if matrix:
            return matrix, u"xlsx:{0}".format(sheet_name)

    try:
        fsize = os.path.getsize(file_path)
    except Exception:
        fsize = -1
    hint = u""
    if fsize == 0:
        hint = u"（文件大小为 0，请重新导出）"
    elif _file_is_xlsx(file_path):
        hint = u"（内容为 xlsx 格式，请改用 .xlsx 扩展名或导出为 UTF-8 CSV）"
    elif last_err and "field larger than field limit" in to_text(last_err).lower():
        hint = u"（单字段过长，请使用新版脚本或 export MYSQL_COMMENT_ENCODING=gb18030）"
    else:
        hint = u"（若为中文 CSV，请 export MYSQL_COMMENT_ENCODING=gb18030）"
    raise RuntimeError(
        u"无法读取 MySQL 注释文件：{0}{1} size={2} last_err={3}；可尝试 {4}".format(
            file_path, hint, fsize, to_text(last_err), DEFAULT_MYSQL_COMMENT_XLSX
        )
    )


def build_csv_comment_index(csv_path):
    """
    将 mysql_comment.csv / mysql_comment.xlsx 构建为内存索引，供 Excel 驱动匹配。
    返回 dict：field_index / table_cn / table_field_keys / stats
    """
    rows_raw, source_desc = _load_comment_source_rows(csv_path)
    _print_u(u"【{0}】MySQL 注释源：{1} ({2})".format(_now(), csv_path, source_desc))

    header = None
    data_start = 0
    for i, row in enumerate(rows_raw):
        if not row:
            continue
        joined = u"".join([_normalize_match_key(c) for c in row])
        if u"用户名" in joined or u"表英文名" in joined or u"字段" in joined:
            header = row
            data_start = i + 1
            break
    if header is None:
        header = rows_raw[0]
        data_start = 1

    idx_user = _pick_header_index(header, [C_USER, u"用户", u"schema", u"database", u"库名"])
    idx_table = _pick_header_index(header, [C_TABLE, u"表名", u"table", u"table_name"])
    idx_table_cn = _pick_header_index(header, [C_TABLE_CN, u"表注释", u"table_comment", u"table comment"])
    idx_field = _pick_header_index(header, [C_FIELD, u"列名", u"column", u"column_name", u"字段"])
    idx_comment = _pick_header_index(header, [C_COMMENT, u"注释", u"comment", u"column_comment"])

    if idx_user is None or idx_table is None:
        raise RuntimeError(
            u"CSV 缺少必要列（需含 用户名、表英文名）。表头={0}".format(u"|".join([to_text(x) for x in header]))
        )
    if idx_field is None or idx_comment is None:
        raise RuntimeError(
            u"CSV 缺少 字段名/字段注释 列。表头={0}".format(u"|".join([to_text(x) for x in header]))
        )

    field_index = {}
    table_cn = {}
    table_field_keys = {}
    raw_count = 0
    for row in rows_raw[data_start:]:
        if not row or not any([_trim_token(x) for x in row]):
            continue
        raw_count += 1
        user = _normalize_match_key(_row_get(row, idx_user))
        table = _normalize_match_key(_row_get(row, idx_table))
        field = _normalize_match_key(_row_get(row, idx_field))
        comment = _trim_token(_row_get(row, idx_comment))
        table_cn_val = _trim_token(_row_get(row, idx_table_cn)) if idx_table_cn is not None else ""
        if not user or not table or not field:
            continue
        k2 = (user, table)
        k3 = (user, table, field)
        if k3 in field_index:
            continue
        field_index[k3] = {
            "user": user,
            "table": table,
            "field": field,
            "comment": comment,
            "table_cn": table_cn_val,
            "user_raw": _row_get(row, idx_user),
            "table_raw": _row_get(row, idx_table),
            "field_raw": _row_get(row, idx_field),
        }
        table_field_keys.setdefault(k2, []).append(field)
        if table_cn_val and k2 not in table_cn:
            table_cn[k2] = table_cn_val
    dedup_count = len(field_index)
    if raw_count > dedup_count:
        _print_u(
            u"【{0}】CSV 索引：原始行={1} 有效字段={2} 去重合并={3} 涉及表={4}".format(
                _now(), raw_count, dedup_count, raw_count - dedup_count, len(table_field_keys)
            )
        )
    else:
        _print_u(
            u"【{0}】CSV 索引：有效字段={1} 涉及表={2}".format(
                _now(), dedup_count, len(table_field_keys)
            )
        )
    return {
        "field_index": field_index,
        "table_cn": table_cn,
        "table_field_keys": table_field_keys,
    }


def _xlsx_cell_value(c_elem, t, shared_strings):
    v_elem = None
    for ch in c_elem:
        if ch.tag.endswith("}v") or ch.tag == "v":
            v_elem = ch
            break
    if v_elem is None or v_elem.text is None:
        return ""
    raw = v_elem.text
    if t == "s":
        try:
            idx = int(raw)
            return shared_strings[idx] if 0 <= idx < len(shared_strings) else ""
        except Exception:
            return raw
    return raw


def _xlsx_load_shared_strings(zf):
    shared = []
    if "xl/sharedStrings.xml" not in zf.namelist():
        return shared
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    for si in root.iter():
        if not (si.tag.endswith("}si") or si.tag == "si"):
            continue
        parts = []
        for t_node in si.iter():
            if t_node.tag.endswith("}t") or t_node.tag == "t":
                if t_node.text:
                    parts.append(t_node.text)
        shared.append(u"".join(parts))
    return shared


def _xlsx_resolve_sheet_path(zf, sheet_name):
    wb = ET.fromstring(zf.read("xl/workbook.xml"))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_id = None
    for sh in wb.findall(".//m:sheet", ns):
        name = sh.get("name") or ""
        if to_text(name).strip() == to_text(sheet_name).strip():
            rel_id = sh.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            break
    if not rel_id:
        names = []
        for sh in wb.findall(".//m:sheet", ns):
            names.append(sh.get("name") or "")
        raise RuntimeError(
            u"未找到工作表 {0!r}，可用：{1}".format(sheet_name, u", ".join(names))
        )
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target = None
    for rel in rels:
        if rel.get("Id") == rel_id:
            target = rel.get("Target")
            break
    if not target:
        raise RuntimeError(u"无法解析工作表路径：{0}".format(sheet_name))
    if not target.startswith("xl/"):
        target = "xl/" + target.lstrip("/")
    return target


def _xlsx_read_sheet_rows_stdlib(xlsx_path, sheet_name):
    with zipfile.ZipFile(xlsx_path, "r") as zf:
        shared = _xlsx_load_shared_strings(zf)
        sheet_xml = _xlsx_resolve_sheet_path(zf, sheet_name)
        root = ET.fromstring(zf.read(sheet_xml))
        cells = {}
        for c in root.iter():
            if not (c.tag.endswith("}c") or c.tag == "c"):
                continue
            ref = c.get("r") or ""
            m = _CELL_REF_RE.match(to_text(ref).upper())
            if not m:
                continue
            col_letters, row_s = m.group(1), m.group(2)
            try:
                rn = int(row_s)
            except Exception:
                continue
            val = _xlsx_cell_value(c, c.get("t") or "", shared)
            cells.setdefault(rn, {})[col_letters] = to_text(val)

    if not cells:
        return []
    max_row = max(cells.keys())
    col_letters_sorted = sorted(
        {cl for rn in cells for cl in cells[rn].keys()},
        key=lambda x: (len(x), x),
    )
    col_index = {cl: i for i, cl in enumerate(col_letters_sorted)}
    max_col = len(col_letters_sorted)
    matrix = []
    for rn in range(1, max_row + 1):
        row = [u""] * max_col
        if rn in cells:
            for cl, val in cells[rn].items():
                if cl in col_index:
                    row[col_index[cl]] = val
        matrix.append(row)
    return matrix


def load_mapping_xlsx(xlsx_path, sheet_name):
    if not os.path.isfile(xlsx_path):
        raise RuntimeError("Excel 文件不存在：{0}".format(xlsx_path))

    matrix = None
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            alt = None
            for sn in wb.sheetnames:
                if to_text(sn).strip() == to_text(sheet_name).strip():
                    alt = sn
                    break
            if alt is None:
                raise RuntimeError(u"未找到工作表 {0!r}，可用：{1}".format(sheet_name, u", ".join(wb.sheetnames)))
            sheet_name = alt
        ws = wb[sheet_name]
        matrix = []
        try:
            max_col = ws.max_column or 1
            max_row = ws.max_row or 1
        except Exception:
            max_col, max_row = 200, 50000
        for row in ws.iter_rows(
            min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True
        ):
            matrix.append([to_text(c) if c is not None else u"" for c in row])
        wb.close()
    except ImportError:
        matrix = _xlsx_read_sheet_rows_stdlib(xlsx_path, sheet_name)
    except Exception as e:
        if "未找到工作表" in to_text(e):
            raise
        matrix = _xlsx_read_sheet_rows_stdlib(xlsx_path, sheet_name)

    if not matrix:
        raise RuntimeError(u"Excel 工作表为空：{0}".format(sheet_name))

    header = None
    data_start = 0
    for i, row in enumerate(matrix):
        if not row:
            continue
        joined = u"".join([_normalize_match_key(c) for c in row])
        if u"二期" in joined and (u"源表" in joined or u"中心库" in joined):
            header = row
            data_start = i + 1
            break
    if header is None:
        header = matrix[0]
        data_start = 1

    idx_center = _pick_header_index(header, [F_CENTER])
    idx_source = _pick_header_index(header, [F_SOURCE, u"二期源表"])
    idx_hive_db = _pick_header_index(
        header, [F_HIVE_DB, u"三期-hive-ods 库", u"三期hive-ods库", u"三期-ods库", u"hive-ods库"]
    )
    idx_hive_table = _pick_header_index(
        header, [F_HIVE_TABLE, u"三期-ods表名", u"ods表英文名", u"三期ods表英文名"]
    )
    idx_field = _pick_header_index(header, [F_FIELD, u"三期-字段名", u"hive字段名"])
    idx_migrate = _pick_header_index(
        header,
        [F_MIGRATE_FLAG, u"是否迁移", u"迁移", u"是否迁移标识", u"迁移标识栏位"],
    )

    missing = []
    if idx_center is None:
        missing.append(F_CENTER)
    if idx_source is None:
        missing.append(F_SOURCE)
    if idx_hive_table is None:
        missing.append(F_HIVE_TABLE)
    if idx_migrate is None:
        missing.append(F_MIGRATE_FLAG)
    if missing:
        raise RuntimeError(
            u"Excel 缺少列：{0}。表头={1}".format(u", ".join(missing), u"|".join([to_text(x) for x in header if x]))
        )

    default_hive_db = _normalize_identifier_token(os.getenv("HIVE_ODS_DATABASE", "") or "")
    if idx_hive_db is None and not default_hive_db:
        _print_u(
            u"【{0}】WARNING: Excel 未找到「{1}」，且未设置环境变量 HIVE_ODS_DATABASE，"
            u"将无法确定 Hive 库名".format(_now(), F_HIVE_DB)
        )

    migrate_rows = []
    skipped_not_migrate = 0
    for row in matrix[data_start:]:
        if not row or not any([_trim_token(x) for x in row]):
            continue
        if not _flag_is_yes(_row_get(row, idx_migrate)):
            skipped_not_migrate += 1
            continue
        center = _normalize_match_key(_row_get(row, idx_center))
        source = _normalize_match_key(_row_get(row, idx_source))
        hive_db = _normalize_identifier_token(_row_get(row, idx_hive_db)) if idx_hive_db is not None else ""
        if not hive_db:
            hive_db = default_hive_db
        hive_table = _normalize_identifier_token(_row_get(row, idx_hive_table))
        field_raw = _row_get(row, idx_field) if idx_field is not None else ""
        field = _normalize_match_key(field_raw)
        if not center or not source or not hive_db or not hive_table:
            continue
        migrate_rows.append(
            {
                "center": center,
                "source": source,
                "hive_db": hive_db,
                "hive_table": hive_table,
                "field": field,
                "field_display": field_raw,
            }
        )

    _print_u(
        u"【{0}】Excel 迁移标识过滤：{1}=是 共 {2} 条，跳过非「是」 {3} 条".format(
            _now(), F_MIGRATE_FLAG, len(migrate_rows), skipped_not_migrate
        )
    )
    if not migrate_rows:
        _print_u(
            u"【{0}】WARNING: 终版页无「迁移标识=是」的有效映射行，后续将无法匹配 CSV".format(_now())
        )

    return migrate_rows, skipped_not_migrate


def _get_or_create_group(groups, hive_db, hive_table, center="", source=""):
    gkey = (hive_db, hive_table)
    if gkey not in groups:
        groups[gkey] = {
            "hive_db": hive_db,
            "hive_table": hive_table,
            "table_comment": u"",
            "center": _normalize_match_key(center),
            "source": _normalize_match_key(source),
            "_fields_map": {},
        }
    else:
        g = groups[gkey]
        if center and not g.get("center"):
            g["center"] = _normalize_match_key(center)
        if source and not g.get("source"):
            g["source"] = _normalize_match_key(source)
    return groups[gkey]


def _set_table_comment(group, table_cn):
    tc = _trim_token(table_cn)
    if not tc:
        return
    old = _trim_token(group.get("table_comment"))
    if not old:
        group["table_comment"] = tc
    elif old != tc:
        group.setdefault("_table_comment_conflict", True)


def _append_field_to_group(group, field_key, field_display, comment, overwrite=False):
    fmap = group["_fields_map"]
    if field_key in fmap and not overwrite:
        return
    fmap[field_key] = {
        "field": field_key,
        "field_display": field_display,
        "comment": comment,
    }


def _hive_target_has_comments(old_table_cmt, cols):
    """Hive 表或任一字段是否已有非空注释。"""
    if _trim_token(old_table_cmt):
        return True
    for meta in (cols or {}).values():
        if _trim_token(meta.get("comment")):
            return True
    return False


def _csv_direct_when_hive_has_comment():
    return _env_truthy("HIVE_COMMENT_CSV_WHEN_HIVE_HAS", True)


def _csv_force_reapply_when_hive_has():
    """
    Hive 已有注释时是否强制用 CSV 覆盖（即使字符串相同也执行 ALTER）。
    默认开启；设 HIVE_COMMENT_FORCE_REAPPLY=0 则仅在不同时更新。
    """
    if (os.getenv("HIVE_COMMENT_FORCE_REAPPLY", "") or "").strip():
        return _env_truthy("HIVE_COMMENT_FORCE_REAPPLY", True)
    return True


def _ensure_group_fields_map(group):
    """终版分组可能仅有 fields 列表，统一为可写的 _fields_map。"""
    if "_fields_map" in group:
        return group["_fields_map"]
    fmap = {}
    for job in group.get("fields") or []:
        fk = job.get("field") or ""
        if fk:
            fmap[fk] = job
    group["_fields_map"] = fmap
    return fmap


def _sync_group_fields_list(group):
    fmap = group.get("_fields_map")
    if fmap is not None:
        group["fields"] = list(fmap.values())


def _expand_group_fields_from_csv(group, csv_index):
    """
    Hive 目标已有注释时：将该二期源表在 mysql_comment.csv 中的全部字段并入更新计划，
    注释一律以 CSV 为准（覆盖 Excel 子集字段列表）。
    """
    center = group.get("center") or ""
    source = group.get("source") or ""
    if not center or not source:
        return 0
    k2 = (center, source)
    field_index = csv_index.get("field_index") or {}
    table_field_keys = csv_index.get("table_field_keys") or {}
    table_cn_map = csv_index.get("table_cn") or {}
    fmap = _ensure_group_fields_map(group)
    before = set(fmap.keys())
    for field_key in table_field_keys.get(k2) or []:
        item = field_index.get((center, source, field_key))
        if not item:
            continue
        _append_field_to_group(
            group,
            field_key,
            item.get("field_raw") or field_key,
            item["comment"],
            overwrite=True,
        )
    added = len(set(fmap.keys()) - before)
    tc = table_cn_map.get(k2) or ""
    if tc:
        group["table_comment"] = _trim_token(tc)
    return added


def _finalize_groups(groups):
    out = []
    for gkey in sorted(groups.keys()):
        g = groups[gkey]
        g["fields"] = list(g.pop("_fields_map", {}).values())
        if g.pop("_table_comment_conflict", False):
            _print_u(
                u"【{0}】WARNING: 表 {1}.{2} 存在多个表中文名，已保留首次非空值".format(
                    _now(), g["hive_db"], g["hive_table"]
                )
            )
        out.append(g)
    return out


def _summarize_duplicate_hive_targets(migrate_rows):
    """
    统计多条 Excel 迁移行指向同一 Hive 表 (库, 表名) 的情况。
    返回 duplicates 列表；merged_row_count = Excel 行数 - 唯一 Hive 表数（合并掉的行数）。
    """
    buckets = {}
    for er in migrate_rows or []:
        gkey = (er.get("hive_db") or "", er.get("hive_table") or "")
        buckets.setdefault(gkey, []).append(
            {
                "center": er.get("center") or "",
                "source": er.get("source") or "",
                "field": er.get("field") or "",
                "field_display": er.get("field_display") or "",
            }
        )
    duplicates = []
    merged_row_count = 0
    for (hive_db, hive_table), rows in sorted(buckets.items()):
        if len(rows) <= 1:
            continue
        merged_row_count += len(rows) - 1
        seen = set()
        unique_mappings = []
        for r in rows:
            sig = (r["center"], r["source"], r["field"])
            if sig in seen:
                continue
            seen.add(sig)
            unique_mappings.append(r)
        duplicates.append(
            {
                "hive_db": hive_db,
                "hive_table": hive_table,
                "excel_row_count": len(rows),
                "unique_mapping_count": len(unique_mappings),
                "mappings": unique_mappings,
            }
        )
    return duplicates, merged_row_count


def _print_duplicate_hive_targets_summary(migrate_rows, hive_group_count, use_excel_field=False):
    """输出 Excel 迁移行合并到同一 Hive 目标表的汇总。"""
    excel_row_count = len(migrate_rows or [])
    duplicates, merged_row_count = _summarize_duplicate_hive_targets(migrate_rows)
    if merged_row_count > 0:
        _print_u(
            u"【{0}】Excel 迁移行合并：共 {1} 行 -> Hive 唯一表 {2} 张（重复映射合并 {3} 行）".format(
                _now(), excel_row_count, hive_group_count, merged_row_count
            )
        )
    if not duplicates:
        return
    dup_excel_rows = sum(d["excel_row_count"] for d in duplicates)
    _print_u(
        u"【{0}】重复映射 Hive 表汇总（共 {1} 张，涉及 Excel 行 {2} 条）".format(
            _now(), len(duplicates), dup_excel_rows
        )
    )
    show_field = use_excel_field
    for i, item in enumerate(duplicates, 1):
        full_table = u"{0}.{1}".format(item["hive_db"], item["hive_table"])
        _print_u(
            u"  [{0}] {1} | Excel行={2} | 源映射={3}".format(
                i,
                full_table,
                item["excel_row_count"],
                item["unique_mapping_count"],
            )
        )
        for m in item["mappings"]:
            line = u"      - {0}={1} | {2}={3}".format(
                F_CENTER, m["center"], F_SOURCE, m["source"]
            )
            if show_field and (m.get("field") or m.get("field_display")):
                fd = m.get("field_display") or m.get("field")
                line += u" | {0}={1}".format(F_FIELD, fd)
            _print_u(line)
    print("=" * 80)


def _group_tables_by_hive_db(groups):
    """
    按 Excel「三期-hive-ods库」分组，同库表连续排列，便于共用 PyHive 连接。
    返回 [(hive_db, [group, ...]), ...]，按库名排序。
    """
    buckets = {}
    for g in groups or []:
        db = _safe_hive_ident(g.get("hive_db")) or "default"
        buckets.setdefault(db, []).append(g)
    out = []
    for db in sorted(buckets.keys()):
        tables = sorted(
            buckets[db],
            key=lambda x: _normalize_match_key(x.get("hive_table")),
        )
        out.append((db, tables))
    return out


def _empty_process_stats():
    return {
        "ok_table": 0,
        "fail_table": 0,
        "ok_col": 0,
        "fail_col": 0,
        "skip_table_same": 0,
        "skip_col_same": 0,
        "skip_missing_col": 0,
        "skip_missing_table": 0,
        "skip_missing_fields": 0,
        "hive_calls": 0,
        "skip_empty_tables": 0,
        "failures": [],
        "missing_tables": [],
    }


def _merge_process_stats(total, part):
    for k in (
        "ok_table",
        "fail_table",
        "ok_col",
        "fail_col",
        "skip_table_same",
        "skip_col_same",
        "skip_missing_col",
        "skip_missing_table",
        "skip_missing_fields",
        "hive_calls",
        "skip_empty_tables",
    ):
        total[k] = total.get(k, 0) + part.get(k, 0)
    total.setdefault("failures", []).extend(part.get("failures") or [])
    total.setdefault("missing_tables", []).extend(part.get("missing_tables") or [])
    return total


def build_grouped_update_plan(csv_index, migrate_rows):
    """
    以 Excel（迁移标识=是）为驱动，在 CSV 索引中查找注释；
    按 (三期-hive-ods库, 三期-ods表英文名) 汇总分组。
    """
    groups = {}
    skipped_no_csv_table = 0
    skipped_no_csv_field = 0
    matched_fields = 0

    field_index = csv_index.get("field_index") or {}
    table_cn_map = csv_index.get("table_cn") or {}
    table_field_keys = csv_index.get("table_field_keys") or {}

    use_excel_field = any(_trim_token(er.get("field")) for er in migrate_rows)

    if use_excel_field:
        for er in migrate_rows:
            center = er["center"]
            source = er["source"]
            field_key = er.get("field") or ""
            if not field_key:
                continue
            k3 = (center, source, field_key)
            item = field_index.get(k3)
            if not item:
                skipped_no_csv_field += 1
                continue
            matched_fields += 1
            group = _get_or_create_group(
                groups, er["hive_db"], er["hive_table"], center, source
            )
            _set_table_comment(group, table_cn_map.get((center, source), item.get("table_cn") or ""))
            fd = er.get("field_display") or item.get("field_raw") or field_key
            _append_field_to_group(group, field_key, fd, item["comment"])
    else:
        table_targets = {}
        for er in migrate_rows:
            k2 = (er["center"], er["source"])
            tgt = (er["hive_db"], er["hive_table"])
            if k2 in table_targets and table_targets[k2] != tgt:
                _print_u(
                    u"【{0}】WARNING: Excel 映射冲突，保留首条：{1}+{2} -> {3}（忽略 {4}）".format(
                        _now(),
                        k2[0],
                        k2[1],
                        u"{0}.{1}".format(table_targets[k2][0], table_targets[k2][1]),
                        u"{0}.{1}".format(tgt[0], tgt[1]),
                    )
                )
            else:
                table_targets[k2] = tgt

        for k2, (hive_db, hive_table) in table_targets.items():
            fields = table_field_keys.get(k2) or []
            if not fields:
                skipped_no_csv_table += 1
                continue
            group = _get_or_create_group(
                groups, hive_db, hive_table, k2[0], k2[1]
            )
            _set_table_comment(group, table_cn_map.get(k2, ""))
            for field_key in fields:
                item = field_index.get((k2[0], k2[1], field_key))
                if not item:
                    continue
                matched_fields += 1
                fd = item.get("field_raw") or field_key
                _append_field_to_group(group, field_key, fd, item["comment"])

    return _finalize_groups(groups), skipped_no_csv_table, skipped_no_csv_field, matched_fields, use_excel_field


def plan_pending_updates_for_group(group, cols, old_table_cmt, hive_has_comments=False):
    """
    对比 Hive 现状，生成待执行的 ALTER 语句列表（不含 USE）。
    hive_has_comments=True 时：目标已有注释的表/字段以 mysql_comment.csv 为准直接更新。
    返回 (alter_stmts, stats_dict)
    """
    stats = {
        "skip_table_same": 0,
        "pending_table": 0,
        "skip_col_same": 0,
        "pending_col": 0,
        "skip_missing_col": 0,
        "csv_direct_table": 0,
        "csv_direct_col": 0,
    }
    alter_stmts = []
    hive_table = group["hive_table"]
    csv_direct = hive_has_comments and _csv_direct_when_hive_has_comment()
    force_overwrite = csv_direct and _csv_force_reapply_when_hive_has()

    new_table_cmt = _trim_token(group.get("table_comment"))
    old_tbl = _trim_token(old_table_cmt)
    if new_table_cmt:
        if force_overwrite:
            stats["csv_direct_table"] = 1
            alter_stmts.append(build_alter_table_comment_stmt(hive_table, new_table_cmt))
            stats["pending_table"] = 1
        elif old_tbl == new_table_cmt:
            stats["skip_table_same"] = 1
        else:
            alter_stmts.append(build_alter_table_comment_stmt(hive_table, new_table_cmt))
            stats["pending_table"] = 1

    for job in group.get("fields") or []:
        col_meta = _resolve_hive_column(cols, job["field"], job["field_display"])
        if not col_meta:
            stats["skip_missing_col"] += 1
            continue
        old_cmt = _trim_token(col_meta.get("comment"))
        new_cmt = _trim_token(job.get("comment"))
        if force_overwrite:
            stats["csv_direct_col"] += 1
            alter_stmts.append(
                build_alter_column_comment_stmt(
                    hive_table, col_meta["name"], col_meta["type"], new_cmt
                )
            )
            stats["pending_col"] += 1
        elif old_cmt == new_cmt:
            stats["skip_col_same"] += 1
            continue
        else:
            alter_stmts.append(
                build_alter_column_comment_stmt(
                    hive_table, col_meta["name"], col_meta["type"], new_cmt
                )
            )
            stats["pending_col"] += 1

    return alter_stmts, stats


def process_table_group(
    group,
    hive_cfg,
    dry_run,
    pyhive_pool=None,
    csv_index=None,
    shared_pyhive_cur=None,
):
    """单表：复用同库共享 PyHive 游标读写元数据与 DDL。"""
    hive_db = group["hive_db"]
    hive_table = group["hive_table"]
    full_table = u"{0}.{1}".format(hive_db, hive_table)
    result = {
        "ok_table": 0,
        "fail_table": 0,
        "ok_col": 0,
        "fail_col": 0,
        "skip_table_same": 0,
        "skip_col_same": 0,
        "skip_missing_col": 0,
        "hive_calls": 0,
        "skipped_empty": False,
        "failures": [],
        "missing_tables": [],
    }

    field_cnt = len(group.get("fields") or [])
    _print_u(u"【{0}】汇总表：{1}（待处理字段={2}）".format(_now(), full_table, field_cnt))
    hive_has_comments = False

    if not pyhive_pool or not pyhive_pool.enabled:
        raise RuntimeError(u"PyHive 会话未启用，无法处理表：{0}".format(full_table))

    pyhive_cur = shared_pyhive_cur
    if pyhive_cur is None:
        _, pyhive_cur = pyhive_pool.acquire(hive_db)

    try:
        try:
            cols, old_table_cmt, table_type, resolved_table = (
                _hive_fetch_table_metadata_pyhive_resolved(
                    pyhive_cur, hive_db, hive_table
                )
            )
            if resolved_table and resolved_table != hive_table:
                group["hive_table"] = resolved_table
                hive_table = resolved_table
                full_table = u"{0}.{1}".format(hive_db, hive_table)
        except Exception as ex:
            if _is_hive_session_down_error(ex):
                _print_u(
                    u"【{0}】WARNING: PyHive 读元数据 session 断开，重连：{1}".format(
                        _now(), full_table
                    )
                )
                _, pyhive_cur = pyhive_pool.reconnect(hive_db)
                cols, old_table_cmt, table_type, resolved_table = (
                    _hive_fetch_table_metadata_pyhive_resolved(
                        pyhive_cur, hive_db, hive_table
                    )
                )
                if resolved_table and resolved_table != group.get("hive_table"):
                    group["hive_table"] = resolved_table
                    hive_table = resolved_table
                    full_table = u"{0}.{1}".format(hive_db, hive_table)
            else:
                raise
        result["hive_calls"] += 1
        if table_type:
            _print_u(
                u"【{0}】Hive 对象类型：{1}.{2} TableType={3}".format(
                    _now(), hive_db, hive_table, table_type
                )
            )
        if _hive_is_virtual_view(table_type, hive_table):
            _print_u(
                u"【{0}】跳过 Hive 视图（仅 VIRTUAL_VIEW/MATERIALIZED_VIEW）：{1}".format(
                    _now(), full_table
                )
            )
            result["skipped_empty"] = True
            return result
        hive_has_comments = _hive_target_has_comments(old_table_cmt, cols)
        if hive_has_comments and csv_index and _csv_direct_when_hive_has_comment():
            expanded = _expand_group_fields_from_csv(group, csv_index)
            if expanded:
                _print_u(
                    u"【{0}】Hive 已有注释，按 CSV 全量字段强制覆盖：{1}（CSV 字段 {2} 条）".format(
                        _now(), full_table, len(group.get("_fields_map") or {})
                    )
                )
            else:
                _print_u(
                    u"【{0}】Hive 已有注释，强制以 mysql_comment.csv 覆盖：{1}（字段 {2} 条）".format(
                        _now(), full_table, len(group.get("_fields_map") or {})
                    )
                )
            _sync_group_fields_list(group)
            field_cnt = len(group.get("fields") or [])
    except Exception as ex:
        err_msg = to_text(ex)
        if _is_hive_table_not_found_error(ex):
            _print_u(
                u"【{0}】跳过（Hive 表不存在）：{1}（字段计划 {2} 条，请核对三期-ods表英文名）".format(
                    _now(), full_table, field_cnt
                )
            )
            result["skip_missing_table"] = 1
            result["skip_missing_fields"] = field_cnt
            result["missing_tables"].append(
                _make_missing_table_skip_record(
                    full_table,
                    group.get("table_comment"),
                    field_cnt,
                    err_msg,
                )
            )
            return result
        err_short = err_msg[:800]
        _print_u(
            u"【{0}】读取元数据失败：{1} err={2}".format(_now(), full_table, err_short)
        )
        result["failures"].append(
            _make_fail_record(
                full_table,
                "metadata",
                u"",
                group.get("table_comment"),
                u"",
                err_msg,
            )
        )
        result["fail_col"] = field_cnt
        if _trim_token(group.get("table_comment")):
            result["fail_table"] = 1
        return result

    alter_stmts, pst = plan_pending_updates_for_group(
        group, cols, old_table_cmt, hive_has_comments=hive_has_comments
    )
    result["skip_table_same"] = pst.get("skip_table_same", 0)
    result["skip_col_same"] = pst.get("skip_col_same", 0)
    result["skip_missing_col"] = pst.get("skip_missing_col", 0)

    if not alter_stmts:
        result["skipped_empty"] = True
        _print_u(
            u"【{0}】跳过执行（无变更）：{1} table_skip={2} col_skip={3} missing_col={4}".format(
                _now(),
                full_table,
                result["skip_table_same"],
                result["skip_col_same"],
                result["skip_missing_col"],
            )
        )
        return result

    ddl_batches = max(1, (len(alter_stmts) + _hive_ddl_batch_size() - 1) // _hive_ddl_batch_size())
    ok_t, fail_t, ok_c, fail_c, ddl_failures = _execute_alter_stmts_chunked(
        hive_db,
        alter_stmts,
        hive_cfg,
        dry_run,
        full_table,
        n_table_stmts=pst.get("pending_table", 0),
        pyhive_cur=pyhive_cur,
        pyhive_pool=pyhive_pool,
    )
    result["hive_calls"] += ddl_batches
    result["ok_table"] = ok_t
    result["fail_table"] = fail_t
    result["ok_col"] = ok_c
    result["fail_col"] = fail_c
    result["failures"].extend(ddl_failures)
    if fail_t or fail_c:
        _print_u(
            u"【{0}】部分失败 {1}：表注释 ok={2} fail={3} 字段注释 ok={4} fail={5}".format(
                _now(), full_table, ok_t, fail_t, ok_c, fail_c
            )
        )
    else:
        _print_u(
            u"【{0}】OK {1}：表注释={2} 字段注释={3} ddl_batches={4}".format(
                _now(), full_table, ok_t, ok_c, ddl_batches
            )
        )

    if pyhive_cur is not None:
        result["pyhive_cur"] = pyhive_cur
    return result


def process_hive_db_batch(hive_db, table_groups, hive_cfg, dry_run, pyhive_pool=None, csv_index=None):
    """
    按「三期-hive-ods库」批量处理：同库只加载一次配置、只建一条 PyHive 连接（若可用）。
    """
    stats = _empty_process_stats()
    table_groups = table_groups or []
    if not table_groups:
        return stats

    db_ident = _safe_hive_ident(hive_db) or "default"
    field_total = sum(len(g.get("fields") or []) for g in table_groups)
    _print_u(u"=" * 60)
    _print_u(
        u"【{0}】开始处理 Hive 库：{1}（本库 {2} 张表，字段计划 {3} 条）".format(
            _now(), db_ident, len(table_groups), field_total
        )
    )

    if not pyhive_pool or not pyhive_pool.enabled:
        raise RuntimeError(u"PyHive 未启用，无法处理库：{0}".format(db_ident))

    _, shared_cur = pyhive_pool.acquire(db_ident)
    _print_u(
        u"【{0}】库 {1} 已建立 PyHive 会话，本库所有表共用此连接".format(_now(), db_ident)
    )

    for idx, group in enumerate(table_groups, 1):
        full_table = u"{0}.{1}".format(db_ident, group.get("hive_table"))
        try:
            pr = process_table_group(
                group,
                hive_cfg,
                dry_run,
                pyhive_pool=pyhive_pool,
                csv_index=csv_index,
                shared_pyhive_cur=shared_cur,
            )
        except Exception as ex:
            err_msg = to_text(ex)
            field_cnt = len(group.get("fields") or [])
            if _is_hive_table_not_found_error(ex):
                _print_u(
                    u"【{0}】跳过（Hive 表不存在）：{1}（字段计划 {2} 条）".format(
                        _now(), full_table, field_cnt
                    )
                )
                pr = {
                    "ok_table": 0,
                    "fail_table": 0,
                    "ok_col": 0,
                    "fail_col": 0,
                    "skip_table_same": 0,
                    "skip_col_same": 0,
                    "skip_missing_col": 0,
                    "skip_missing_table": 1,
                    "skip_missing_fields": field_cnt,
                    "hive_calls": 0,
                    "skipped_empty": False,
                    "failures": [],
                    "missing_tables": [
                        _make_missing_table_skip_record(
                            full_table,
                            group.get("table_comment"),
                            field_cnt,
                            err_msg,
                        )
                    ],
                }
            else:
                _print_u(
                    u"【{0}】处理表异常（继续下一张）：{1} err={2}".format(
                        _now(), full_table, err_msg
                    )
                )
                pr = {
                    "ok_table": 0,
                    "fail_table": 1 if _trim_token(group.get("table_comment")) else 0,
                    "ok_col": 0,
                    "fail_col": field_cnt,
                    "skip_table_same": 0,
                    "skip_col_same": 0,
                    "skip_missing_col": 0,
                    "skip_missing_table": 0,
                    "skip_missing_fields": 0,
                    "hive_calls": 0,
                    "skipped_empty": False,
                    "failures": [
                        _make_fail_record(
                            full_table, "metadata", u"", group.get("table_comment"), u"", err_msg
                        )
                    ],
                    "missing_tables": [],
                }
        if pr.get("pyhive_cur") is not None:
            shared_cur = pr.get("pyhive_cur")
        _merge_process_stats(stats, pr)
        if pr.get("skipped_empty"):
            stats["skip_empty_tables"] += 1

    _print_u(
        u"【{0}】库 {1} 处理完成：表注释 ok={2} fail={3} 字段 ok={4} fail={5} "
        u"无变更跳过={6} 张 表不存在跳过={7} 张".format(
            _now(),
            db_ident,
            stats.get("ok_table", 0),
            stats.get("fail_table", 0),
            stats.get("ok_col", 0),
            stats.get("fail_col", 0),
            stats.get("skip_empty_tables", 0),
            stats.get("skip_missing_table", 0),
        )
    )
    if pyhive_pool is not None and pyhive_pool.enabled:
        pyhive_pool.release(db_ident)
    _print_u(u"=" * 60)
    return stats


def _count_planned_comment_updates(groups):
    """匹配完成后统计计划涉及的表注释、字段注释条数。"""
    table_with_cn = 0
    field_total = 0
    for g in groups or []:
        if _trim_token(g.get("table_comment")):
            table_with_cn += 1
        field_total += len(g.get("fields") or [])
    return len(groups or []), table_with_cn, field_total


def _print_comment_update_summary(
    title,
    dry_run,
    ok_table,
    ok_col,
    fail_table,
    fail_col,
    skip_table_same,
    skip_col_same,
    skip_missing_col,
    hive_groups=0,
    skip_empty_tables=0,
    skip_missing_table=0,
    skip_missing_fields=0,
):
    """输出表注释 / 字段注释更新数量汇总。"""
    tag = u"【预览】" if dry_run else u""
    _print_u(
        u"【{0}】{1}{2}".format(_now(), tag, title)
    )
    _print_u(
        u"  表注释：更新 {0} 个，失败 {1} 个，已相同跳过 {2} 个，表不存在跳过 {3} 个".format(
            ok_table, fail_table, skip_table_same, skip_missing_table
        )
    )
    _print_u(
        u"  字段注释：更新 {0} 个，失败 {1} 个，已相同跳过 {2} 个，Hive无列跳过 {3} 个，"
        u"表不存在跳过 {4} 个".format(
            ok_col, fail_col, skip_col_same, skip_missing_col, skip_missing_fields
        )
    )
    if hive_groups:
        _print_u(
            u"  Hive 表组：共 {0} 张，整表无变更跳过 {1} 张，表不存在跳过 {2} 张".format(
                hive_groups, skip_empty_tables, skip_missing_table
            )
        )


def _resolve_hive_column(describe_cols, field_key, field_display):
    if field_key in describe_cols:
        return describe_cols[field_key]
    fd = _normalize_match_key(field_display)
    if fd in describe_cols:
        return describe_cols[fd]
    return None


def _resolve_mysql_comment_path(cli_csv_arg):
    """解析 MySQL 注释文件：优先入参/环境变量，否则 /tmp/lgbs 下 csv 或 xlsx。"""
    if cli_csv_arg:
        return os.path.abspath(to_text(cli_csv_arg))
    env_p = (os.getenv("MYSQL_COMMENT_CSV", "") or "").strip()
    if env_p:
        return os.path.abspath(to_text(env_p))
    if os.path.isfile(DEFAULT_MYSQL_COMMENT_CSV):
        return DEFAULT_MYSQL_COMMENT_CSV
    if os.path.isfile(DEFAULT_MYSQL_COMMENT_XLSX):
        return DEFAULT_MYSQL_COMMENT_XLSX
    return DEFAULT_MYSQL_COMMENT_CSV


def _resolve_input_paths(cli_args):
    """
    解析 CSV / Excel 路径：CLI 入参优先，其次环境变量，最后 /tmp/lgbs 默认文件。
    """
    xlsx_path = (os.getenv("MIGRATION_XLSX", "") or os.getenv("XLSX_PATH", "") or "").strip()
    if not xlsx_path:
        xlsx_path = DEFAULT_MIGRATION_XLSX

    pos = [a for a in (cli_args or []) if a and not a.startswith("-")]
    csv_path = _resolve_mysql_comment_path(pos[0] if len(pos) >= 1 else None)
    if len(pos) >= 2:
        xlsx_path = pos[1]

    xlsx_path = os.path.abspath(to_text(xlsx_path))
    return csv_path, xlsx_path


def main():
    dry_run = False
    args = [a for a in sys.argv[1:] if a]
    if "--dry-run" in args:
        dry_run = True
        args = [a for a in args if a != "--dry-run"]
    if os.getenv("DRY_RUN", "").strip().lower() in ("1", "true", "yes"):
        dry_run = True

    csv_path, xlsx_path = _resolve_input_paths(args)
    if not os.path.isfile(csv_path) and not os.path.isfile(DEFAULT_MYSQL_COMMENT_XLSX):
        _print_u(u"【{0}】ERROR: MySQL 注释文件不存在：{1}".format(_now(), csv_path))
        _print_u(
            u"用法：python hive-comment-update-from-mysql.py [--dry-run]\n"
            u"  或：python hive-comment-update-from-mysql.py <mysql_comment.csv|xlsx> <mapping.xlsx> [--dry-run]\n"
            u"默认：{0} 或 {1}\n       {2}".format(
                DEFAULT_MYSQL_COMMENT_CSV, DEFAULT_MYSQL_COMMENT_XLSX, DEFAULT_MIGRATION_XLSX
            )
        )
        sys.exit(2)
    if not os.path.isfile(xlsx_path):
        _print_u(u"【{0}】ERROR: Excel 不存在：{1}".format(_now(), xlsx_path))
        sys.exit(2)
    sheet_name = (os.getenv("XLSX_SHEET_NAME", "") or u"终版").strip() or u"终版"

    print("=" * 80)
    _print_u(u"【{0}】MySQL 注释 -> Hive 表/字段注释更新 开始".format(_now()))
    _print_u(u"【{0}】CSV={1}".format(_now(), csv_path))
    _print_u(u"【{0}】Excel={1} sheet={2} dry_run={3}".format(_now(), xlsx_path, sheet_name, dry_run))
    print("=" * 80)

    csv_index = build_csv_comment_index(csv_path)
    migrate_rows, skipped_not_migrate = load_mapping_xlsx(xlsx_path, sheet_name)
    groups, skipped_no_csv_table, skipped_no_csv_field, matched_fields, use_excel_field = (
        build_grouped_update_plan(csv_index, migrate_rows)
    )
    total_fields = sum(len(g.get("fields") or []) for g in groups)
    mode = u"Excel驱动-按字段" if use_excel_field else u"Excel驱动-按表"

    _print_u(
        u"【{0}】匹配完成：mode={1} csv_index_fields={2} excel_migrate_yes={3} "
        u"excel_skip_not_yes={4} matched_fields={5} hive_groups={6} pending_fields={7} "
        u"skip_no_csv_table={8} skip_no_csv_field={9}".format(
            _now(),
            mode,
            len(csv_index.get("field_index") or {}),
            len(migrate_rows),
            skipped_not_migrate,
            matched_fields,
            len(groups),
            total_fields,
            skipped_no_csv_table,
            skipped_no_csv_field,
        )
    )
    _print_duplicate_hive_targets_summary(migrate_rows, len(groups), use_excel_field)

    if not groups:
        _print_u(u"【{0}】无待更新任务，退出".format(_now()))
        return

    hive_group_cnt, plan_table_cn, plan_field_cnt = _count_planned_comment_updates(groups)
    db_batches = _group_tables_by_hive_db(groups)
    _print_u(
        u"【{0}】待更新计划：Hive 表 {1} 张（含表中文名 {2} 张），字段注释 {3} 条".format(
            _now(), hive_group_cnt, plan_table_cn, plan_field_cnt
        )
    )
    _print_u(
        u"【{0}】按 Excel「{1}」分组：共 {2} 个库，同库连续处理、共用连接".format(
            _now(), F_HIVE_DB, len(db_batches)
        )
    )
    for db, tbls in db_batches:
        _print_u(u"  - {0}：{1} 张表".format(db, len(tbls)))

    try:
        get_hive_config("default")
    except RuntimeError as ex:
        _print_u(u"【{0}】ERROR: {1}".format(_now(), to_text(ex)))
        sys.exit(1)

    if not _hpc_pyhive_import_available():
        err = _hpc_last_pyhive_import_error() or u"import pyhive 失败"
        _print_u(u"【{0}】ERROR: PyHive 不可用：{1}".format(_now(), err))
        sys.exit(1)

    pyhive_pool = _PyHiveSessionPool()
    try:
        _lgbs_path = _hpc_lgbs_script_candidates()[0]
        _test_conn = _hpc_open_pyhive_connection(get_hive_config("default"), database="default")
        _close_pyhive(_test_conn, None)
        pyhive_pool.enabled = True
        _print_u(
            u"【{0}】Hive 连接：PyHive（{1}，按库复用会话）batch_size={2}".format(
                _now(), _lgbs_path, _hive_ddl_batch_size()
            )
        )
    except Exception as e:
        _print_u(u"【{0}】ERROR: PyHive 建连失败：{1}".format(_now(), to_text(e)))
        sys.exit(1)

    run_stats = _empty_process_stats()
    all_failures = []
    all_missing_tables = []

    try:
        for hive_db, table_groups in db_batches:
            hive_cfg = get_hive_config(hive_db)
            db_stats = process_hive_db_batch(
                hive_db,
                table_groups,
                hive_cfg,
                dry_run,
                pyhive_pool=pyhive_pool,
                csv_index=csv_index,
            )
            _merge_process_stats(run_stats, db_stats)
            all_failures.extend(db_stats.get("failures") or [])
            all_missing_tables.extend(db_stats.get("missing_tables") or [])
    finally:
        pyhive_pool.close_all()

    ok_table = run_stats.get("ok_table", 0)
    fail_table = run_stats.get("fail_table", 0)
    skip_table_same = run_stats.get("skip_table_same", 0)
    ok_col = run_stats.get("ok_col", 0)
    fail_col = run_stats.get("fail_col", 0)
    skip_col_same = run_stats.get("skip_col_same", 0)
    skip_missing_col = run_stats.get("skip_missing_col", 0)
    skip_empty_tables = run_stats.get("skip_empty_tables", 0)
    skip_missing_table = run_stats.get("skip_missing_table", 0)
    skip_missing_fields = run_stats.get("skip_missing_fields", 0)
    total_hive_calls = run_stats.get("hive_calls", 0)

    print("=" * 80)
    _print_comment_update_summary(
        u"注释更新完成汇总",
        dry_run,
        ok_table,
        ok_col,
        fail_table,
        fail_col,
        skip_table_same,
        skip_col_same,
        skip_missing_col,
        hive_groups=len(groups),
        skip_empty_tables=skip_empty_tables,
        skip_missing_table=skip_missing_table,
        skip_missing_fields=skip_missing_fields,
    )
    _print_u(
        u"【{0}】执行明细：hive_calls≈{1} dry_run={2}".format(
            _now(), total_hive_calls, dry_run
        )
    )
    _print_missing_tables_summary(all_missing_tables)
    _print_failures_summary(all_failures)
    if not all_failures and not all_missing_tables:
        print("=" * 80)

    strict_exit = _env_truthy("HIVE_COMMENT_STRICT_EXIT", True)
    has_success = (ok_table + ok_col) > 0
    if strict_exit and (fail_table > 0 or fail_col > 0 or all_failures):
        sys.exit(1)
    if not strict_exit and not has_success and (fail_table > 0 or fail_col > 0 or all_failures):
        sys.exit(1)
    if not strict_exit and (fail_table > 0 or fail_col > 0):
        _print_u(
            u"【{0}】WARNING: 存在失败项但 HIVE_COMMENT_STRICT_EXIT=0，作业返回成功（exit 0）".format(
                _now()
            )
        )


if __name__ == "__main__":
    main()
